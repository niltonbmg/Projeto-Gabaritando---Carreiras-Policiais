#!/usr/bin/env python3
"""Gera o Gabarito Comentado do Simulado 01 – CFOA PMBA 2026 no padrão
visual da revisão de Formosa do Rio Preto (capa amarela, gabarito rápido,
análise alternativa por alternativa e DICA GABARITANDO).

Saídas: saidas/gabarito_comentado_cfoa_pmba_simulado01.{md,docx}
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import importlib
import sys

SIM = sys.argv[1] if len(sys.argv) > 1 else "01"
_dados = importlib.import_module(f"dados_cfoa_pmba_simulado{SIM}")
Q = _dados.Q

BASE = Path(__file__).resolve().parent.parent
SAIDAS = BASE / "saidas"
LOGO = SAIDAS / "logo_gabaritando.png"
MARCA = SAIDAS / "marca_dagua_gabaritando.png"
NOME = f"gabarito_comentado_cfoa_pmba_simulado{SIM}"
RODAPE = f"Simulado {SIM} – CFOA PMBA 2026 • Gabaritando • Página "

AMARELO, LARANJA, CINZA = "F5B800", "E08A00", "F2F2F2"
VERDE, VERMELHO, PRETO = RGBColor(0x1B, 0x8A, 0x2F), RGBColor(0xC0, 0x1F, 0x1F), RGBColor(0x1F, 0x1F, 0x1F)
INVERTIDAS = getattr(_dados, "INVERTIDAS", {29, 33, 43, 57, 78, 80})  # comando pede a INCORRETA
PREFIXO = re.compile(r"^(CORRETA|INCORRETA \(gabarito\)|EM DESACORDO \(gabarito\)|GABARITO DA BANCA)\.?\s*")


def tema(n):
    if len(Q[n]) > 5:
        return Q[n][5]
    return re.split(r"[:.;(]", Q[n][3])[0].strip()


def rotulo(n, letra):
    certa = letra == Q[n][1]
    if n in INVERTIDAS:
        return ("CERTA (é a incorreta)", VERDE) if certa else ("ERRADA (afirmação verdadeira)", VERMELHO)
    return ("CERTA", VERDE) if certa else ("ERRADA", VERMELHO)


def comentario(n, letra):
    return PREFIXO.sub("", Q[n][2][letra])


# ---------------- Markdown ----------------
def gerar_md(destino):
    ln = [f"# Gabarito Comentado – Simulado {SIM} – CFOA PMBA 2026", ""]
    disc = None
    for n in sorted(Q):
        d, gab, _, fund, obs = Q[n][:5]
        if d != disc:
            disc = d
            ln += [f"## {d.upper()}", ""]
        ln += [f"**Questão {n:02d} • {tema(n)} — GABARITO: {gab}**", ""]
        for letra in "ABCDE":
            ln.append(f"- {letra.lower()}) **{rotulo(n, letra)[0]}** – {comentario(n, letra)}")
        ln.append(f"- 💡 **DICA GABARITANDO:** {fund}" + (f" ⚠️ {obs}" if obs else ""))
        ln.append("")
    destino.write_text("\n".join(ln))


# ---------------- helpers docx ----------------
def sombrear(el, cor):
    pr = el.get_or_add_tcPr() if hasattr(el, "get_or_add_tcPr") else el.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor)
    pr.append(shd)


def borda_esquerda(p, cor="1F1F1F", tam=24):
    bdr = parse_xml(
        '<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:left w:val="single" w:sz="{tam}" w:space="6" w:color="{cor}"/></w:pBdr>')
    p._p.get_or_add_pPr().append(bdr)


def run(p, texto, bold=False, italic=False, size=10, cor=PRETO):
    r = p.add_run(texto)
    r.bold, r.italic = bold, italic
    r.font.size = Pt(size)
    r.font.color.rgb = cor
    return r


def cabecalho_rodape(doc):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.add_run().add_picture(str(LOGO), width=Cm(2.6))
    # marca d'água no cabeçalho (imagem atrás do texto, centralizada)
    r = hp.add_run()
    r.add_picture(str(MARCA), width=Cm(15))
    inline = r._r.find(qn("w:drawing")).find(qn("wp:inline"))
    ext = inline.find(qn("wp:extent"))
    cx, cy = ext.get("cx"), ext.get("cy")
    graphic = inline.find(qn("a:graphic"))
    anchor = parse_xml(
        '<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" behindDoc="1" '
        'locked="0" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:align>center</wp:align></wp:positionH>'
        '<wp:positionV relativeFrom="page"><wp:align>center</wp:align></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
        '<wp:docPr id="99" name="MarcaDagua"/><wp:cNvGraphicFramePr/></wp:anchor>')
    anchor.append(graphic)
    inline.getparent().replace(inline, anchor)
    # linha amarela sob o cabeçalho
    bdr = parse_xml(
        '<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:bottom w:val="single" w:sz="8" w:space="4" w:color="{AMARELO}"/></w:pBdr>')
    hp._p.get_or_add_pPr().append(bdr)

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(fp, RODAPE, size=8, cor=RGBColor(0x80, 0x80, 0x80))
    fp._p.append(parse_xml(
        '<w:fldSimple xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'w:instr=" PAGE "><w:r><w:rPr><w:sz w:val="16"/><w:color w:val="808080"/></w:rPr>'
        '<w:t>1</w:t></w:r></w:fldSimple>'))


def capa(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(LOGO), width=Cm(9))
    doc.add_paragraph()
    t = doc.add_table(rows=1, cols=1)
    c = t.cell(0, 0)
    sombrear(c._tc, AMARELO)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    run(p, f"SIMULADO {SIM}", bold=True, size=28)
    p = c.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    run(p, "CONCURSO CFOA PMBA 2026", bold=True, size=12)
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, "OFICIAL AUXILIAR (QOA) – TIPO A", bold=True, size=13)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, "Gabarito comentado alternativa por alternativa", italic=True, size=11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, "Q01–Q80  •  80 questões objetivas", size=9)
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, "VEM GABARITAR COM A GENTE", bold=True, size=10, cor=RGBColor(0xE0, 0x8A, 0x00))
    doc.add_page_break()


def sumario_e_gabarito(doc):
    p = doc.add_paragraph()
    run(p, "SUMÁRIO", bold=True, size=14)
    p = doc.add_paragraph()
    p._p.append(parse_xml(
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:fldChar w:fldCharType="begin" w:dirty="true"/>'
        '<w:instrText xml:space="preserve"> TOC \\o "1-2" \\h \\z \\u </w:instrText>'
        '<w:fldChar w:fldCharType="separate"/>'
        '<w:t>Clique com o botão direito e escolha "Atualizar campo" (F9) para gerar o sumário.</w:t>'
        '<w:fldChar w:fldCharType="end"/></w:r>'))
    doc.add_page_break()

    p = doc.add_paragraph()
    run(p, "GABARITO RÁPIDO", bold=True, size=14)
    for ini in (1, 21, 41, 61):
        tab = doc.add_table(rows=2, cols=20)
        tab.style = "Table Grid"
        for i in range(20):
            n = ini + i
            h, v = tab.cell(0, i), tab.cell(1, i)
            sombrear(h._tc, "000000")
            for cel, txt, cor in ((h, f"Q{n:02d}", RGBColor(0xFF, 0xFF, 0xFF)),
                                  (v, Q[n][1] + ("*" if Q[n][4] else ""), PRETO)):
                pp = cel.paragraphs[0]
                pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run(pp, txt, bold=True, size=7, cor=cor)
        doc.add_paragraph()
    p = doc.add_paragraph()
    run(p, "* Questões com observação (ver comentário): " +
        ", ".join(f"Q{n:02d}" for n in sorted(Q) if Q[n][4]) + ".",
        italic=True, size=8, cor=RGBColor(0xE0, 0x8A, 0x00))


def questao(doc, n):
    d, gab, _, fund, obs = Q[n][:5]
    h = doc.add_heading(level=2)
    run(h, f"Questão {n:02d} • {tema(n)}", bold=True, size=11)
    h.paragraph_format.space_before = Pt(14)

    p = doc.add_paragraph()
    run(p, "GABARITO: ", bold=True, size=10)
    r = run(p, gab, bold=True, size=10)
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW

    p = doc.add_paragraph()
    run(p, "Análise alternativa por alternativa", bold=True, size=9)
    for letra in "ABCDE":
        lab, cor = rotulo(n, letra)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        run(p, f"{letra.lower()}) {lab}", bold=True, size=9, cor=cor)
        run(p, f" – {comentario(n, letra)}", size=9)
    if obs:
        p = doc.add_paragraph()
        run(p, "ATENÇÃO – ", bold=True, size=9, cor=RGBColor(0xE0, 0x8A, 0x00))
        run(p, obs, size=9)
    p = doc.add_paragraph()
    sombrear(p._p, CINZA)
    borda_esquerda(p)
    run(p, "💡 DICA GABARITANDO: ", bold=True, size=9)
    run(p, fund, size=9)


def gerar_docx(destino):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    cabecalho_rodape(doc)
    capa(doc)
    sumario_e_gabarito(doc)
    disc = None
    for n in sorted(Q):
        if Q[n][0] != disc:
            disc = Q[n][0]
            doc.add_page_break()
            h = doc.add_heading(level=1)
            run(h, disc.upper(), bold=True, size=14)
        questao(doc, n)
    doc.save(destino)


def main():
    assert sorted(Q) == list(range(1, 81))
    for n, (_, gab, alts, *_) in Q.items():
        assert set(alts) == set("ABCDE") and PREFIXO.match(alts[gab]), n
    SAIDAS.mkdir(exist_ok=True)
    gerar_md(SAIDAS / f"{NOME}.md")
    gerar_docx(SAIDAS / f"{NOME}.docx")
    print("OK: 80 questões")


if __name__ == "__main__":
    main()
