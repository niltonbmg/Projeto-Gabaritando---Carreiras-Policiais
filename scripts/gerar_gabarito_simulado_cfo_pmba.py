#!/usr/bin/env python3
"""
Gera o Gabarito Comentado dos simulados CFO PMBA (Oficial QOPM).

Uso: python3 scripts/gerar_gabarito_simulado_cfo_pmba.py [06|07]   (padrão: 06)
Os dados de cada simulado ficam em scripts/dados_gabarito_simuladoNN_cfo_pmba.py.

Identidade visual do Projeto Gabaritando (a mesma do caderno do simulado):
  - marca d'água com o "G" do Gabaritando (assets/gabaritando/marca_dagua_g.png)
  - rodapé com "OFICIAL (QOPM)" + logo GABARITANDO (assets/gabaritando/logo_rodape.png)

Saídas:
  saidas/gabarito_comentado_simuladoNN_cfo_pmba.docx
  saidas/gabarito_comentado_simuladoNN_cfo_pmba.pdf  (via LibreOffice, se disponível)
"""
import importlib
import subprocess
import sys
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

NUM = sys.argv[1] if len(sys.argv) > 1 else "06"
_dados = importlib.import_module(f"dados_gabarito_simulado{NUM}_cfo_pmba")
Q, DISCURSIVA = _dados.Q, _dados.DISCURSIVA

BASE = Path(__file__).resolve().parent.parent
ASSETS = BASE / "assets" / "gabaritando"
SAIDA = BASE / "saidas" / f"gabarito_comentado_simulado{NUM}_cfo_pmba.docx"

# Paleta da marca (azul-marinho + dourado do logo "Projeto Gabaritando")
NAVY = "1B2A4A"
OURO = "B8862B"
OURO_CLARO = "F6EBD3"
CINZA = "1F2933"
VERDE = "1E7B34"
VERM = "9B1C1C"

A4_W, A4_H = Cm(21.0), Cm(29.7)


# ---------------------------------------------------------------- utilitários
def _shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _run(p, text, bold=False, size=10.5, color=CINZA, italic=False):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    return r


def _para(doc_or_cell, text="", bold=False, size=10.5, color=CINZA, align=None,
          before=0, after=4, italic=False):
    p = doc_or_cell.add_paragraph()
    if text:
        _run(p, text, bold, size, color, italic)
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def _inline_para_anchor(run, largura, altura, pos_x=0, pos_y=0, rel="page"):
    """Converte a imagem inline do run em âncora atrás do texto (marca d'água/fundo)."""
    inline = run._r.find(qn("w:drawing")).find(qn("wp:inline"))
    graphic = inline.find(qn("a:graphic"))
    docpr = inline.find(qn("wp:docPr"))
    anchor = parse_xml(
        '<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" behindDoc="1" '
        'locked="1" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="{rel}"><wp:posOffset>{int(pos_x)}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="{rel}"><wp:posOffset>{int(pos_y)}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{int(largura)}" cy="{int(altura)}"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/></wp:anchor>')
    anchor.append(docpr)
    cnv = OxmlElement("wp:cNvGraphicFramePr")
    anchor.append(cnv)
    anchor.append(graphic)
    inline.getparent().replace(inline, anchor)


def _campo(run, instr):
    for tipo, txt in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText" if txt == instr else "w:t")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)


def _configurar_cabecalho_rodape(section):
    """Marca d'água (G do Gabaritando) no cabeçalho + rodapé com logo GABARITANDO."""
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    r = hp.add_run()
    lado = Cm(15)
    r.add_picture(str(ASSETS / "marca_dagua_g.png"), width=lado, height=lado)
    _inline_para_anchor(r, lado, lado, (A4_W - lado) / 2, (A4_H - lado) / 2)

    fp = section.footer.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(5.0))
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(17.0))
    _run(fp, "OFICIAL (QOPM)\t", size=9, color=CINZA)
    fp.add_run().add_picture(str(ASSETS / "logo_rodape.png"), height=Cm(1.3))
    _run(fp, "\tPágina ", size=9, color=CINZA)
    _campo(_run(fp, "", size=9, color=CINZA), "PAGE")


def _barra_titulo(doc, texto, cor=NAVY, size=13):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    _shade(c, cor)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _run(p, texto, bold=True, size=size, color="FFFFFF")
    _para(doc, after=2)


# ---------------------------------------------------------------- seções
def capa(doc):
    _para(doc, after=30)
    p = _para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    p.add_run().add_picture(str(ASSETS / "logo_projeto.png"), width=Cm(9))
    _para(doc, "VEM GABARITAR COM A GENTE", True, 11, OURO, WD_ALIGN_PARAGRAPH.CENTER, after=40)
    _para(doc, "GABARITO COMENTADO", True, 32, NAVY, WD_ALIGN_PARAGRAPH.CENTER, after=4)
    _para(doc, f"SIMULADO {NUM}", True, 26, OURO, WD_ALIGN_PARAGRAPH.CENTER, after=4)
    _para(doc, "Prova Objetiva e Discursiva", False, 14, CINZA, WD_ALIGN_PARAGRAPH.CENTER, after=30)
    _para(doc, "PMBA – Curso de Formação de Oficiais", True, 15, NAVY, WD_ALIGN_PARAGRAPH.CENTER, after=2)
    _para(doc, "Cargo: Oficial (QOPM)" + (" – Tipo A" if NUM == "06" else ""), False, 12, CINZA, WD_ALIGN_PARAGRAPH.CENTER, after=40)
    _para(doc, "“O crescimento mora no desconforto.”", True, 12, NAVY,
          WD_ALIGN_PARAGRAPH.CENTER, after=4, italic=True)
    _para(doc, "@projeto.gabaritando  ·  #seremosirmaosdefarda", False, 10, CINZA,
          WD_ALIGN_PARAGRAPH.CENTER)


def grade_respostas(doc):
    _barra_titulo(doc, "GABARITO OFICIAL – QUADRO DE RESPOSTAS")
    linhas, cols = 8, 10
    t = doc.add_table(rows=linhas * 2, cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, item in enumerate(Q):
        bloco, col = divmod(i, cols)
        c_num, c_gab = t.cell(bloco * 2, col), t.cell(bloco * 2 + 1, col)
        _shade(c_num, NAVY)
        _shade(c_gab, OURO_CLARO)
        for c, txt, cor in ((c_num, f"{item['n']:02d}", "FFFFFF"), (c_gab, item["gab"], NAVY)):
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(p, txt, bold=True, size=10, color=cor)
    _para(doc, after=6)

    # Distribuição por disciplina
    _barra_titulo(doc, "DISTRIBUIÇÃO DAS QUESTÕES", cor=OURO, size=11)
    cont = Counter(item["disc"] for item in Q)
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for c, txt in zip(t.rows[0].cells, ("Disciplina", "Questões", "Quantidade")):
        _shade(c, NAVY)
        _run(c.paragraphs[0], txt, True, 10, "FFFFFF")
    for disc in dict.fromkeys(item["disc"] for item in Q):
        nums = [item["n"] for item in Q if item["disc"] == disc]
        row = t.add_row().cells
        _run(row[0].paragraphs[0], disc, True, 10, NAVY)
        _run(row[1].paragraphs[0], f"{min(nums):02d} a {max(nums):02d}", size=10)
        _run(row[2].paragraphs[0], str(cont[disc]), size=10)

    ressalvas = [item for item in Q if item["obs"]]
    if ressalvas:
        _para(doc, after=4)
        _para(doc, "Questões com observação (possível recurso/ressalva): "
              + ", ".join(f"{i['n']:02d}" for i in ressalvas) + ".", True, 10, VERM)


def questao(doc, item):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c1, c2 = t.rows[0].cells
    t.autofit = False
    for c, w in ((c1, Cm(12.5)), (c2, Cm(4.5))):
        c.width = w
    t.columns[0].width, t.columns[1].width = Cm(12.5), Cm(4.5)
    _shade(c1, NAVY)
    _shade(c2, OURO)
    p = c1.paragraphs[0]
    _run(p, f"QUESTÃO {item['n']:02d}  ", True, 12, "FFFFFF")
    _run(p, f"{item['disc']} · {item['tema']}", False, 9.5, "FFFFFF")
    p = c2.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, f"GABARITO: {item['gab']}", True, 12, "FFFFFF")

    p = _para(doc, before=4, after=4)
    _run(p, "Comentário: ", True, 10.5, NAVY)
    _run(p, item["com"])

    _para(doc, "Análise das alternativas", True, 10, OURO, after=2)
    for letra, txt in item["alts"].items():
        p = _para(doc, after=1)
        p.paragraph_format.left_indent = Cm(0.4)
        certo = letra == item["gab"]
        _run(p, f"({letra}) ", True, 10, VERDE if certo else VERM)
        _run(p, txt, bold=certo, size=10)

    t = doc.add_table(rows=1, cols=1)
    c = t.cell(0, 0)
    _shade(c, OURO_CLARO)
    p = c.paragraphs[0]
    _run(p, "🎯 Dica Gabaritando: ", True, 10, NAVY)
    _run(p, item["dica"], size=10, italic=True)
    if item["obs"]:
        p = c.add_paragraph()
        _run(p, "⚠ Observação: ", True, 9.5, VERM)
        _run(p, item["obs"], size=9.5)
    _para(doc, after=8)


def discursiva(doc):
    _barra_titulo(doc, "PROVA DISCURSIVA – ESPELHO DE CORREÇÃO")
    p = _para(doc, after=6)
    _run(p, "Tema: ", True, 11, NAVY)
    _run(p, f"“{DISCURSIVA['tema']}”", True, 11, OURO)
    _para(doc, "Texto dissertativo-argumentativo, entre 20 e 30 linhas.", italic=True, after=8)
    for titulo, txt in DISCURSIVA["estrutura"]:
        _para(doc, titulo, True, 10.5, NAVY, after=1)
        _para(doc, txt, after=6)
    _para(doc, "Critérios de avaliação", True, 11, OURO, before=4, after=2)
    for crit in DISCURSIVA["criterios"]:
        p = _para(doc, after=1)
        p.paragraph_format.left_indent = Cm(0.4)
        _run(p, "▸ ", True, 10, OURO)
        _run(p, crit, size=10)


# ---------------------------------------------------------------- principal
def gerar_docx(destino: Path):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.font.size = Pt(10.5)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = A4_W, A4_H
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.6)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.header_distance, sec.footer_distance = Cm(0), Cm(0.8)
    _configurar_cabecalho_rodape(sec)

    capa(doc)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    grade_respostas(doc)

    disc_atual = None
    for item in Q:
        if item["disc"] != disc_atual:
            disc_atual = item["disc"]
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            _barra_titulo(doc, disc_atual.upper(), cor=OURO, size=14)
        questao(doc, item)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    discursiva(doc)

    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destino)
    return destino


def gerar_pdf(docx: Path):
    try:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir",
                        str(docx.parent), str(docx)], check=True, capture_output=True, timeout=300)
    except (FileNotFoundError, subprocess.CalledProcessError, subprocess.TimeoutExpired) as e:
        print(f"PDF não gerado: {e}")
        return None
    return docx.with_suffix(".pdf")


if __name__ == "__main__":
    assert len(Q) == 80 and [i["n"] for i in Q] == list(range(1, 81))
    d = gerar_docx(SAIDA)
    print("DOCX:", d)
    print("PDF:", gerar_pdf(d))
