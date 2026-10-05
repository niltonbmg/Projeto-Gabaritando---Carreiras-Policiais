#!/usr/bin/env python3
"""
Gera o Gabarito Comentado – PND 2025 História – PV_1 (TIPO 01),
alternativa por alternativa, a partir das análises do manual.

Saídas:
  saidas/gabarito_comentado_pnd_historia.md
  saidas/gabarito_comentado_pnd_historia.docx

Se downloads/questoes.json existir, o texto de cada alternativa é incluído
antes do comentário; caso contrário, apenas letra + veredito + comentário.
"""
import ast
import json
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
SAIDAS = BASE / "saidas"


def carregar_analises():
    """Lê ANALISES e CORRECOES_ALT do script do manual sem executá-lo
    (o manual exige downloads/questoes.json na importação)."""
    src = (BASE / "scripts/gerar_manual_pnd_historia.py").read_text()
    valores = {}
    for node in ast.parse(src).body:
        if isinstance(node, ast.Assign) and len(node.targets) == 1:
            nome = getattr(node.targets[0], "id", None)
            if nome in ("ANALISES", "CORRECOES_ALT"):
                valores[nome] = ast.literal_eval(node.value)
    return valores["ANALISES"], valores["CORRECOES_ALT"]


def carregar_textos(correcoes):
    arq = BASE / "downloads/questoes.json"
    if not arq.exists():
        return {}
    textos = {q["numero"]: dict(q["alternativas"]) for q in json.loads(arq.read_text())}
    for num, alts in correcoes.items():
        textos.setdefault(num, {}).update(alts)
    return textos


def secao(num):
    return "Formação Geral Docente" if num <= 30 else "Componente Específico – História"


def gerar_md(analises, textos, destino):
    ln = [
        "# Gabarito Comentado – PND 2025 História (Caderno PV_1 • TIPO 01)",
        "",
        "Comentário alternativa por alternativa das 80 questões objetivas.",
        "Gabaritos baseados na análise do Projeto Gabaritando; confira com o gabarito oficial do INEP.",
        "",
        "## Quadro-resumo",
        "",
        "| Questão | Gabarito | Questão | Gabarito | Questão | Gabarito | Questão | Gabarito |",
        "|:-:|:-:|:-:|:-:|:-:|:-:|:-:|:-:|",
    ]
    nums = sorted(analises)
    for i in range(0, 20):
        linha = []
        for col in range(4):
            n = nums[i + 20 * col]
            linha += [f"{n:02d}", f"**{analises[n]['gabarito']}**"]
        ln.append("| " + " | ".join(linha) + " |")
    ln.append("")

    sec_atual = None
    for n in nums:
        a = analises[n]
        if secao(n) != sec_atual:
            sec_atual = secao(n)
            ln += [f"## {sec_atual}", ""]
        ln += [
            f"### Questão {n:02d} — Gabarito: {a['gabarito']}",
            "",
            f"*Tema:* {a['tema']} · *Nível:* {a['nivel']}",
            "",
        ]
        for letra, (veredito, comentario) in sorted(a["analise_alternativas"].items()):
            marca = "✅" if veredito.startswith("CORRETA") else "❌"
            texto = textos.get(n, {}).get(letra)
            if texto:
                ln.append(f"- **{letra})** {texto}  ")
                ln.append(f"  {marca} **{veredito}** — {comentario}")
            else:
                ln.append(f"- **{letra})** {marca} **{veredito}** — {comentario}")
        ln += ["", f"> **Dica:** {a['dica_estrategica']}", ""]
    destino.write_text("\n".join(ln))


def gerar_docx(analises, textos, destino):
    from docx import Document
    from docx.shared import Pt, RGBColor

    verde, vermelho, azul = RGBColor(0x1B, 0x6E, 0x2F), RGBColor(0x7A, 0x1F, 0x1F), RGBColor(0x0B, 0x3D, 0x91)
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading("Gabarito Comentado – PND 2025 História", 0)
    doc.add_paragraph("Caderno PV_1 • TIPO 01 — comentário alternativa por alternativa das 80 questões objetivas.")

    doc.add_heading("Quadro-resumo", 1)
    nums = sorted(analises)
    tab = doc.add_table(rows=21, cols=8)
    tab.style = "Light Grid Accent 1"
    for col in range(4):
        tab.cell(0, 2 * col).text, tab.cell(0, 2 * col + 1).text = "Questão", "Gab."
        for i in range(20):
            n = nums[i + 20 * col]
            tab.cell(i + 1, 2 * col).text = f"{n:02d}"
            tab.cell(i + 1, 2 * col + 1).text = analises[n]["gabarito"]

    sec_atual = None
    for n in nums:
        a = analises[n]
        if secao(n) != sec_atual:
            sec_atual = secao(n)
            doc.add_page_break()
            doc.add_heading(sec_atual, 1)
        h = doc.add_heading(f"Questão {n:02d} — Gabarito: {a['gabarito']}", 2)
        h.runs[0].font.color.rgb = azul
        p = doc.add_paragraph()
        p.add_run(f"Tema: {a['tema']} · Nível: {a['nivel']}").italic = True
        for letra, (veredito, comentario) in sorted(a["analise_alternativas"].items()):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"{letra}) ").bold = True
            texto = textos.get(n, {}).get(letra)
            if texto:
                p.add_run(texto + "\n")
            r = p.add_run(veredito)
            r.bold = True
            r.font.color.rgb = verde if veredito.startswith("CORRETA") else vermelho
            p.add_run(f" — {comentario}")
        p = doc.add_paragraph()
        p.add_run("Dica: ").bold = True
        p.add_run(a["dica_estrategica"])
    doc.save(destino)


def main():
    analises, correcoes = carregar_analises()
    textos = carregar_textos(correcoes)
    SAIDAS.mkdir(exist_ok=True)
    gerar_md(analises, textos, SAIDAS / "gabarito_comentado_pnd_historia.md")
    gerar_docx(analises, textos, SAIDAS / "gabarito_comentado_pnd_historia.docx")
    print(f"Gabarito comentado gerado: {len(analises)} questões "
          f"({'com' if textos else 'sem'} texto das alternativas).")


if __name__ == "__main__":
    main()
