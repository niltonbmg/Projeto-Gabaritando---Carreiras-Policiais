#!/usr/bin/env python3
"""
Diagrama as apostilas de Ciências Naturais – Soldado CBMBA (Markdown → DOCX + PDF).

Uso: python3 scripts/gerar_apostila_sd_cbmba.py 01 [02 ...]   (ou "todas")
Fonte:  saidas/sd-cbmba/ciencias-naturais/NN_*.md
Saídas: saidas/sd-cbmba/ciencias-naturais/pdf/NN_*.docx e .pdf

Identidade visual do Projeto Gabaritando (a mesma dos gabaritos CFO PMBA):
  - capa com logo, azul-marinho + dourado
  - marca d'água com o "G" do Gabaritando
  - rodapé com "SOLDADO (CBMBA)" + logo GABARITANDO + número da página
"""
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from docx.oxml import OxmlElement, parse_xml

BASE = Path(__file__).resolve().parent.parent
ASSETS = BASE / "assets" / "gabaritando"
FONTE = BASE / "saidas" / "sd-cbmba" / "ciencias-naturais"
DESTINO = FONTE / "pdf"
AZUL_CLARO = "E8EDF5"

# Paleta da marca (azul-marinho + dourado do logo "Projeto Gabaritando")
NAVY = "1B2A4A"
OURO = "B8862B"
OURO_CLARO = "F6EBD3"
CINZA = "1F2933"
VERDE = "1E7B34"
VERM = "9B1C1C"

A4_W, A4_H = Cm(21.0), Cm(29.7)


def _shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _run(p, text, bold=False, size=10.5, color=CINZA, italic=False):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    return r


def _para(doc_or_cell, text="", bold=False, size=10.5, color=CINZA, align=None,
          before=0, after=4, italic=False):
    p = doc_or_cell.add_paragraph()
    if text:
        _run(p, text, bold, size, color, italic)
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def _inline_para_anchor(run, largura, altura, pos_x=0, pos_y=0, rel="page"):
    """Converte a imagem inline do run em âncora atrás do texto (marca d'água/fundo)."""
    inline = run._r.find(qn("w:drawing")).find(qn("wp:inline"))
    graphic = inline.find(qn("a:graphic"))
    docpr = inline.find(qn("wp:docPr"))
    anchor = parse_xml(
        '<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" behindDoc="1" '
        'locked="1" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="{rel}"><wp:posOffset>{int(pos_x)}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="{rel}"><wp:posOffset>{int(pos_y)}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{int(largura)}" cy="{int(altura)}"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/></wp:anchor>')
    anchor.append(docpr)
    cnv = OxmlElement("wp:cNvGraphicFramePr")
    anchor.append(cnv)
    anchor.append(graphic)
    inline.getparent().replace(inline, anchor)


def _campo(run, instr):
    for tipo, txt in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText" if txt == instr else "w:t")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)




# ---------------------------------------------------------------- utilitários
def _inline(p, texto, size=10.5, color=CINZA, bold=False, italic=False):
    """Escreve texto com **negrito** e *itálico* do Markdown."""
    for parte in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", texto):
        if not parte:
            continue
        if parte.startswith("**"):
            _run(p, parte[2:-2], True, size, NAVY if color == CINZA else color, italic)
        elif parte.startswith("*"):
            _run(p, parte[1:-1], bold, size, color, True)
        else:
            _run(p, parte, bold, size, color, italic)


def _rodape(section):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    r = section.header.paragraphs[0].add_run()
    lado = Cm(15)
    r.add_picture(str(ASSETS / "marca_dagua_g.png"), width=lado, height=lado)
    _inline_para_anchor(r, lado, lado, (A4_W - lado) / 2, (A4_H - lado) / 2)

    fp = section.footer.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(5.0))
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(17.0))
    _run(fp, "SOLDADO (CBMBA)\t", size=9)
    fp.add_run().add_picture(str(ASSETS / "logo_rodape.png"), height=Cm(1.3))
    _run(fp, "\tPágina ", size=9)
    _campo(_run(fp, "", size=9), "PAGE")


def _barra(doc, texto, cor=NAVY, size=13):
    _para(doc, after=4)
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    _shade(c, cor)
    _run(c.paragraphs[0], texto, True, size, "FFFFFF")
    _para(doc, after=2)


def _caixa(doc, cor_fundo, linhas):
    t = doc.add_table(rows=1, cols=1)
    c = t.cell(0, 0)
    _shade(c, cor_fundo)
    p = c.paragraphs[0]
    for i, (texto, kw) in enumerate(linhas):
        if i:
            p = c.add_paragraph()
        _inline(p, texto, **kw)
    _para(doc, after=6)


def _tabela(doc, linhas, cabecalho):
    ncol = max(len(l) for l in linhas)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, cols in enumerate(linhas):
        cells = t.add_row().cells
        cols = cols + [""] * (ncol - len(cols))
        for j, txt in enumerate(cols):
            p = cells[j].paragraphs[0]
            if i == 0 and cabecalho:
                _shade(cells[j], NAVY)
                _inline(p, txt, 9.5, "FFFFFF", bold=True)
            else:
                if i % 2 == (0 if cabecalho else 1):
                    _shade(cells[j], AZUL_CLARO)
                _inline(p, txt, 9.5, NAVY if j == 0 else CINZA, bold=j == 0)
    _para(doc, after=6)


# ---------------------------------------------------------------- leitura do .md
def ler(md: Path):
    linhas = md.read_text(encoding="utf-8").splitlines()
    titulos = [l[2:].strip() for l in linhas if l.startswith("# ")]
    meta = dict(re.findall(r"^\*\*(.+?):\*\*\s*(.+)$", "\n".join(linhas), re.M))
    corpo = linhas[next(i for i, l in enumerate(linhas) if l.startswith("## ")):]
    return titulos[1], meta, corpo


def capa(doc, titulo, meta):
    _para(doc, after=30)
    p = _para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    p.add_run().add_picture(str(ASSETS / "logo_projeto.png"), width=Cm(9))
    _para(doc, "VEM GABARITAR COM A GENTE", True, 11, OURO, WD_ALIGN_PARAGRAPH.CENTER, after=40)
    _para(doc, "APOSTILA", True, 30, NAVY, WD_ALIGN_PARAGRAPH.CENTER, after=4)
    _para(doc, meta.get("Disciplina", "").upper(), True, 18, OURO, WD_ALIGN_PARAGRAPH.CENTER, after=14)
    _para(doc, titulo, True, 15, NAVY, WD_ALIGN_PARAGRAPH.CENTER, after=30)
    _para(doc, "Corpo de Bombeiros Militar da Bahia", True, 14, NAVY, WD_ALIGN_PARAGRAPH.CENTER, after=2)
    _para(doc, f"Cargo: Soldado · Item {meta.get('Item do edital', '')} do edital · "
          f"{meta.get('Ano-base', '')}", False, 12, CINZA, WD_ALIGN_PARAGRAPH.CENTER, after=40)
    _para(doc, "“O crescimento mora no desconforto.”", True, 12, NAVY,
          WD_ALIGN_PARAGRAPH.CENTER, after=4, italic=True)
    _para(doc, "@projeto.gabaritando  ·  #seremosirmaosdefarda", False, 10, CINZA,
          WD_ALIGN_PARAGRAPH.CENTER)


def corpo(doc, linhas):
    secao, tabela, i = "", [], 0

    def fecha_tabela():
        if tabela:
            _tabela(doc, tabela, any(k in secao for k in ("Resumo", "Quadro")))
            tabela.clear()

    while i < len(linhas):
        l = linhas[i].rstrip()
        i += 1
        eh_linha_tabela = (" – " in l and not l.startswith(("–", "**", "*", "#", "_"))
                           and not re.match(r"^[a-e]\) ", l))
        if not eh_linha_tabela:
            fecha_tabela()
        if not l.strip() or l == "---" or l.startswith("_Elaborado"):
            continue
        if eh_linha_tabela:
            tabela.append([c.strip() for c in l.split(" – ")])
        elif l.startswith("## "):
            secao = l[3:]
            if "Questões" in secao or "Desenvolvimento" in secao:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            _barra(doc, secao.upper())
        elif l.startswith("### "):
            _para(doc, l[4:], True, 12.5, OURO, before=8, after=3)
        elif l.startswith("#### "):
            _para(doc, l[5:], True, 11, NAVY, before=4, after=2)
        elif l.startswith("– "):
            p = _para(doc, after=2)
            p.paragraph_format.left_indent = Cm(0.5)
            _run(p, "▸ ", True, 10, OURO)
            _inline(p, l[2:])
        elif re.match(r"^\*\*\d+\.\*\*", l):
            m = re.match(r"^\*\*(\d+)\.\*\*\s*(.*)$", l)
            t = doc.add_table(rows=1, cols=1)
            _shade(t.cell(0, 0), NAVY)
            _run(t.cell(0, 0).paragraphs[0], f"QUESTÃO {int(m.group(1)):02d}", True, 11, "FFFFFF")
            p = _para(doc, before=4, after=3)
            _inline(p, m.group(2))
        elif re.match(r"^[a-e]\) ", l):
            p = _para(doc, after=1)
            p.paragraph_format.left_indent = Cm(0.6)
            _run(p, f"({l[0].upper()}) ", True, 10, NAVY)
            _inline(p, l[3:], 10)
        elif l.startswith("**Gabarito:"):
            gab = re.sub(r"\*\*Gabarito:\s*(.+?)\*\*", r"\1", l)
            com = linhas[i].strip() if i < len(linhas) and linhas[i].startswith("Comentário:") else ""
            if com:
                i += 1
            _caixa(doc, OURO_CLARO, [
                (f"✔ GABARITO: {gab}", dict(size=10.5, color=VERDE, bold=True)),
                (com.replace("Comentário:", "**Comentário:**", 1), dict(size=10))])
        elif l.startswith("**Exemplo"):
            _para(doc, l.strip("*"), True, 10.5, OURO, before=6, after=2)
        elif secao.startswith("✨"):
            _caixa(doc, OURO_CLARO, [(l, dict(size=11, italic=True))])
        else:
            p = _para(doc, after=4)
            _inline(p, l)
    fecha_tabela()


# ---------------------------------------------------------------- principal
def gerar(md: Path):
    titulo, meta, linhas = ler(md)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.font.size = Pt(10.5)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = A4_W, A4_H
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.6)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.header_distance, sec.footer_distance = Cm(0), Cm(0.8)
    _rodape(sec)

    capa(doc, titulo, meta)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    corpo(doc, linhas)

    DESTINO.mkdir(parents=True, exist_ok=True)
    docx = DESTINO / f"{md.stem}.docx"
    doc.save(docx)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(DESTINO),
                    str(docx)], check=True, capture_output=True, timeout=300)
    return docx, docx.with_suffix(".pdf")


if __name__ == "__main__":
    args = sys.argv[1:] or ["01"]
    arquivos = sorted(FONTE.glob("*.md")) if args == ["todas"] else \
        [next(FONTE.glob(f"{n}_*.md")) for n in args]
    for md in arquivos:
        print(*gerar(md))
