#!/usr/bin/env python3
"""
Gera apostila no padrão Projeto Gabaritando (faixa preta no cabeçalho, marca
d'água do "G", sumário, apresentação, assuntos e revisão geral).

Uso:
  python3 scripts/gerar_apostila.py conteudo/tecnico-enfermagem-ses-to

A pasta de conteúdo deve ter:
  config.json   – textos da capa, banner e nome do arquivo de saída
  00_apresentacao.md, 01_*.md, ..., 99_revisao_geral.md

Marcação dos .md (simples, uma linha por bloco):
  # Título          -> título de assunto (nova página, entra no sumário)
  ## Subtítulo      -> seção do assunto (1. Resumo teórico, 2. Pontos...)
  ### Subtítulo     -> subseção
  **linha**         -> parágrafo inteiro em negrito (ex.: QUESTÃO 1, Gabarito)
  - item            -> marcador "• item"
  | a | b |         -> linhas consecutivas viram tabela (1ª linha = cabeçalho)
  > texto           -> caixa de destaque
  texto           -> parágrafo justificado; **trecho** em negrito no meio
Saídas: saidas/<arquivo>.docx e saidas/<arquivo>.pdf (via LibreOffice)
"""
import json
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent.parent

AZUL = RGBColor(0x00, 0x00, 0x00)  # preto (identidade Gabaritando)
DOURADO = RGBColor(0xE5, 0xA8, 0x00)  # amarelo Gabaritando
TEXTO = RGBColor(0x00, 0x00, 0x00)


def _shade(cell, hex_cor):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_cor)
    tc_pr.append(shd)


def _campo(run, instr):
    for tipo, texto in (("begin", None), (None, instr), ("separate", None), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = texto
        run._r.append(el)


def _cabecalho_rodape(section, banner: Path, marca: Path):
    section.header_distance = Cm(0)
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    run.add_picture(str(banner), width=Cm(7.0))

    # marca d'água: imagem ancorada atrás do texto, centralizada na página
    run = p.add_run()
    run.add_picture(str(marca), width=Cm(15.5))
    inline = run._r.find(qn("w:drawing"))[0]
    cx, cy = inline.find(qn("wp:extent")).get("cx"), inline.find(qn("wp:extent")).get("cy")
    graphic = inline.find("{http://schemas.openxmlformats.org/drawingml/2006/main}graphic")
    anchor = parse_xml(
        '<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" behindDoc="1" '
        'locked="1" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:align>center</wp:align></wp:positionH>'
        '<wp:positionV relativeFrom="page"><wp:align>center</wp:align></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
        '<wp:docPr id="999" name="marca-dagua"/><wp:cNvGraphicFramePr/></wp:anchor>')
    anchor.append(graphic)
    drawing = run._r.find(qn("w:drawing"))
    drawing.remove(inline)
    drawing.append(anchor)

    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _campo(fp.add_run(), "PAGE")


class Apostila:
    def __init__(self, cfg):
        self.cfg = cfg
        self.doc = Document()
        st = self.doc.styles["Normal"]
        st.font.name = "Calibri"
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.font.size = Pt(11.5)
        st.paragraph_format.line_spacing = 1.5
        st.paragraph_format.space_after = Pt(6)
        for nivel, tam in ((1, 16), (2, 14), (3, 12)):
            h = self.doc.styles[f"Heading {nivel}"]
            h.font.name = "Calibri"
            h.font.size = Pt(tam)
            h.font.bold = True
            h.font.color.rgb = AZUL
            h.paragraph_format.space_before = Pt(14 if nivel < 3 else 8)
            h.paragraph_format.space_after = Pt(6)
            rfonts = h.element.get_or_add_rPr().get_or_add_rFonts()
            for a in ("w:ascii", "w:hAnsi", "w:cs"):
                rfonts.set(qn(a), "Calibri")
            for a in ("w:asciiTheme", "w:hAnsiTheme"):
                rfonts.attrib.pop(qn(a), None)
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.top_margin, sec.bottom_margin = Cm(3.6), Cm(2.0)
        sec.left_margin = sec.right_margin = Cm(2.5)

    # ---------- blocos ----------
    def par(self, texto, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=None,
            color=None, italic=False):
        p = self.doc.add_paragraph()
        p.alignment = align
        partes = re.split(r"(\*\*[^*]+\*\*)", texto)
        for parte in partes:
            if not parte:
                continue
            negrito = parte.startswith("**") and parte.endswith("**")
            r = p.add_run(parte[2:-2] if negrito else parte)
            r.bold = bold or negrito
            r.italic = italic
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
        return p

    def tabela(self, linhas):
        rows = [[c.strip() for c in l.strip().strip("|").split("|")] for l in linhas
                if not re.fullmatch(r"\|?[\s:\-|]+\|?", l.strip())]
        t = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                c = t.cell(i, j)
                c.text = ""
                p = c.paragraphs[0]
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(val.replace("**", ""))
                r.font.size = Pt(10)
                if i == 0:
                    r.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xC4, 0x00)
                    _shade(c, "000000")
                elif i % 2 == 0:
                    _shade(c, "FFF8DC")
        self.doc.add_paragraph()

    def destaque(self, texto):
        t = self.doc.add_table(rows=1, cols=1)
        c = t.cell(0, 0)
        _shade(c, "FFE58A")
        p = c.paragraphs[0]
        p.paragraph_format.line_spacing = 1.2
        for parte in re.split(r"(\*\*[^*]+\*\*)", texto):
            if parte:
                neg = parte.startswith("**")
                p.add_run(parte.strip("*") if neg else parte).bold = neg
        self.doc.add_paragraph()

    def quebra(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---------- partes ----------
    def capa(self):
        c = self.cfg
        for _ in range(4):
            self.doc.add_paragraph()
        C = WD_ALIGN_PARAGRAPH.CENTER
        for i, linha in enumerate(c["orgao"]):
            p = self.par(linha, bold=True, align=C, size=(15 if i == 0 else 20), color=AZUL)
            for r in p.runs:
                r.font.name = "Calibri"
                r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Calibri")
            p.paragraph_format.space_after = Pt(2)
        self.par(c["subtitulo"], align=C)
        self.par(c["banca"], align=C, italic=True)
        self.par(c["ano"], align=C)
        self.doc.add_paragraph()
        self.par("APOSTILA PREPARATÓRIA COMPLETA", bold=True, align=C, size=22, color=DOURADO)
        self.doc.add_paragraph()
        for linha in c["titulo_capa"]:
            self.par(linha, bold=True, align=C, size=24, color=AZUL)
        self.doc.add_paragraph()
        self.par("Teoria completa • Pontos que mais caem • Macetes • Questões comentadas • "
                 "Revisão por capítulo e geral", align=C, italic=True, size=9.5)
        self.quebra()

    def sumario(self, titulos):
        self.par("SUMÁRIO", bold=True, size=15, color=AZUL, align=WD_ALIGN_PARAGRAPH.LEFT)
        for t, recuo in titulos:
            p = self.par(t, align=WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.left_indent = Cm(0.3 if recuo else 0)
            p.paragraph_format.space_after = Pt(2)
        self.quebra()

    def markdown(self, texto, primeira_pagina_titulo=None):
        linhas = texto.splitlines()
        i = 0
        primeiro_h1 = True
        while i < len(linhas):
            l = linhas[i].rstrip()
            if not l.strip():
                i += 1
                continue
            if l.startswith("|"):
                bloco = []
                while i < len(linhas) and linhas[i].startswith("|"):
                    bloco.append(linhas[i])
                    i += 1
                self.tabela(bloco)
                continue
            if l.startswith("# "):
                if not primeiro_h1:
                    self.quebra()
                primeiro_h1 = False
                self.doc.add_heading(l[2:], level=1)
            elif l.startswith("## "):
                h = self.doc.add_heading(l[3:], level=2)
                pbdr = OxmlElement("w:pBdr")
                left = OxmlElement("w:left")
                for k, v in (("w:val", "single"), ("w:sz", "36"), ("w:space", "6"), ("w:color", "FFC400")):
                    left.set(qn(k), v)
                pbdr.append(left)
                h._p.get_or_add_pPr().append(pbdr)
            elif l.startswith("### "):
                self.doc.add_heading(l[4:], level=3)
            elif l.startswith("- "):
                p = self.par("• " + l[2:])
                p.paragraph_format.left_indent = Cm(0.4)
            elif l.startswith("> "):
                self.destaque(l[2:])
            elif l.startswith("**") and l.endswith("**") and l.count("**") == 2:
                self.par(l[2:-2], bold=True)
            else:
                self.par(l)
            i += 1


def main(pasta):
    pasta = (BASE / pasta).resolve() if not Path(pasta).is_absolute() else Path(pasta)
    cfg = json.loads((pasta / "config.json").read_text(encoding="utf-8"))
    arquivos = sorted(pasta.glob("*.md"))

    titulos = []
    for f in arquivos:
        for l in f.read_text(encoding="utf-8").splitlines():
            if l.startswith("# ") and not l.startswith("# CONHECIMENTOS"):
                t = l[2:].strip()
                titulos.append((t[0] + t[1:].lower() if t.isupper() else t, False))
            elif l.startswith("## Assunto"):
                titulos.append((l[3:].strip(), True))

    ap = Apostila(cfg)
    _cabecalho_rodape(ap.doc.sections[0], BASE / cfg["banner"], BASE / cfg["marca_dagua"])
    ap.capa()
    ap.sumario(titulos)
    for idx, f in enumerate(arquivos):
        if idx:
            ap.quebra()
        ap.markdown(f.read_text(encoding="utf-8"))

    saidas = BASE / "saidas"
    saidas.mkdir(exist_ok=True)
    docx = saidas / f"{cfg['arquivo']}.docx"
    ap.doc.save(docx)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(saidas), str(docx)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    print(docx)
    print(docx.with_suffix(".pdf"))


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "conteudo/tecnico-enfermagem-ses-to")
